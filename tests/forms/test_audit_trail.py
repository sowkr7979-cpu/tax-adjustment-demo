"""감사추적 Excel — 소득처분 단일 소스 정합성 테스트."""
import os
import tempfile
from datetime import date

from openpyxl import load_workbook

from src.forms.audit_trail import generate_audit_trail
from src.forms.summary_rows import adjustment_rows
from src.utils.models import TaxAdjustmentResult


def _make_result() -> TaxAdjustmentResult:
    r = TaxAdjustmentResult(
        fiscal_year_start=date(2025, 1, 1),
        fiscal_year_end=date(2025, 12, 31),
        is_sme=True,
    )
    r.donation_excess = 7_000_000
    r.depreciation_excess = 3_000_000
    r.pension_deduction = 2_000_000   # 차감조정(△유보)
    return r


def test_donation_disposition_is_other_outflow():
    """기부금 한도초과 소득처분은 '기타사외유출' (과거 audit_trail '유보' 오류 회귀 방지)."""
    add_items, _ = adjustment_rows(_make_result())
    don = [it for it in add_items if "기부금" in it[1]]
    assert don, "기부금 행이 있어야 함"
    assert don[0][4] == "기타사외유출"


def test_audit_trail_sheet2_uses_single_source():
    """감사추적 시트2가 summary_rows 처분과 일치하고, 차감조정은 음수로 표기."""
    r = _make_result()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out = f.name
    try:
        generate_audit_trail([], r, out, "(주)테스트", "model", "0.1.0", date(2025, 12, 31))
        wb = load_workbook(out)
        ws = wb["세무조정계산근거"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        by_name = {row[1]: row for row in rows}
        # 기부금: 기타사외유출, 양수
        assert by_name["기부금 한도초과·비지정"][4] == "기타사외유출"
        assert by_name["기부금 한도초과·비지정"][2] == 7_000_000
        # 퇴직연금 부담금: 차감조정 → 음수, △유보
        assert by_name["퇴직연금 부담금"][2] == -2_000_000
        assert by_name["퇴직연금 부담금"][4] == "△유보"
        # 0원 항목은 시트에 없어야 함
        assert "대손충당금 한도초과" not in by_name
    finally:
        os.unlink(out)


def test_audit_trail_has_reserve_sheet():
    """자본금적립금(을) 시트가 생성되고 유보 잔액이 반영된다."""
    r = _make_result()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out = f.name
    try:
        generate_audit_trail(
            [], r, out, "(주)테스트", "model", "0.1.0", date(2025, 12, 31),
            prior_reserves=[{"code": "대손충당금 한도초과", "amount": 1_000_000, "disposition": "유보"}],
            depr_denial_end=9_000_000,
        )
        wb = load_workbook(out)
        assert "자본금적립금(을)" in wb.sheetnames
        ws = wb["자본금적립금(을)"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        names = [row[0] for row in rows]
        assert "감가상각 부인누계" in names
        # 감가상각 기말은 denial_end(9,000,000)에 정합
        dep = next(row for row in rows if row[0] == "감가상각 부인누계")
        assert dep[4] == 9_000_000
    finally:
        os.unlink(out)
