"""감사추적 Excel 생성."""
from __future__ import annotations
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.utils.models import LLMAnalysisResult, TaxAdjustmentResult
from src.forms.summary_rows import adjustment_rows


def _header_style(ws, row: int, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="366092")
    font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def generate_audit_trail(
    llm_results: list[LLMAnalysisResult],
    tax_result: TaxAdjustmentResult,
    output_path: str | Path,
    company_name: str,
    model_name: str,
    rule_engine_version: str,
    law_ref_date: date,
) -> None:
    wb = Workbook()

    # ── 시트1: LLM 분석 결과 ────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "LLM분석결과"
    headers1 = [
        "전표번호", "분개라인번호", "이슈가능여부", "이슈코드",
        "세무조정유형", "해당금액", "신뢰도", "검토필요", "검토사유",
        "관련법령", "목표서식", "LLM프롬프트버전", "모델명", "법령기준일",
    ]
    _header_style(ws1, 1, headers1)
    for r in llm_results:
        legal = "; ".join(
            f"{b.get('법령', '')} {b.get('내용', '')}"
            for b in r.legal_basis_candidates
        )
        forms = "; ".join(r.target_form_candidates)
        ws1.append([
            r.journal_id, r.line_id,
            "가능" if r.issue_possible else "없음",
            r.tax_issue_code.value,
            r.tax_adjustment_type, r.affected_amount,
            round(r.stage2_confidence, 3),
            "Y" if r.review_required else "N",
            r.review_reason, legal, forms,
            "v0.1", model_name, str(law_ref_date),
        ])

    # ── 시트2: 세무조정 계산 근거 ────────────────────────────────────────────
    # 소득금액조정합계표(별지15호)와 동일 소스(summary_rows.adjustment_rows) 사용 —
    # 항목·소득처분을 두 곳에서 따로 관리하지 않는다 (정합성 단일 소스).
    ws2 = wb.create_sheet("세무조정계산근거")
    headers2 = ["구분", "항목", "금액(원)", "근거법령", "소득처분"]
    _header_style(ws2, 1, headers2)
    add_items, deduct_items = adjustment_rows(tax_result)
    # 가산조정은 금액 양수, 차감조정(손금산입·익금불산입)은 음수로 표기. 0원 항목은 제외.
    for gubun, name, amount, basis, disposition in add_items:
        if amount:
            ws2.append([gubun, name, amount, basis, disposition])
    for gubun, name, amount, basis, disposition in deduct_items:
        if amount:
            ws2.append([gubun, name, -amount, basis, disposition])

    # ── 시트3: 세액 계산 근거 ────────────────────────────────────────────────
    ws3 = wb.create_sheet("세액계산근거")
    headers3 = ["항목", "금액(원)"]
    _header_style(ws3, 1, headers3)
    rows3 = [
        ("각사업연도 소득금액", tax_result.business_income),
        ("과세표준", tax_result.tax_base),
        ("산출세액", tax_result.gross_tax),
        ("최저한세_적용_감면합계", tax_result.tax_credits_min_subject),
        ("최저한세액", tax_result.min_tax),
        ("배제된 감면", tax_result.excluded_credits),
        ("최저한세_미적용_감면합계", tax_result.tax_credits_post_min),
        ("가산세", tax_result.surtax),
        ("기납부세액", tax_result.prepaid_tax),
        ("차감납부세액", tax_result.final_tax_due),
    ]
    for row in rows3:
        ws3.append(list(row))

    # 열 너비 자동 조정
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    wb.save(output_path)
