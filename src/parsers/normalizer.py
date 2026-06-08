"""회계자료 표준화 파이프라인 — 업로드 파일을 파싱 가능한 표 형태로 정규화.

회계 프로그램(더존 Smart A·WEHAGO 등) 내보내기 파일의 현실적인 문제를 처리한다:
  1. .xls 확장자인데 실제로는 HTML 저장 파일 (구버전 더존 내보내기)
  2. cp949/euc-kr 인코딩 HTML·CSV
  3. 제목·회사명·빈 행이 헤더 위에 섞여 있어 컬럼 인식 실패
  4. 음수 표기 다양성: (1,234) / △1,234 / ▲1,234 / -1,234
  5. 합계·소계행과 실제 계정행 혼재
  6. 회사·프로그램마다 다른 계정과목명 (동의어)

모든 읽기는 read_any_table()을 통과시키고, 무엇을 어떻게 인식했는지
meta dict로 기록해 UI에서 진단 리포트로 보여준다 (표준화 투명성).
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import pandas as pd

# ── ① 파일 포맷 스니핑 ────────────────────────────────────────────────────────

_OLE2_MAGIC = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"   # 진짜 XLS (BIFF)
_ZIP_MAGIC = b"PK\x03\x04"                            # 진짜 XLSX (OOXML)


def sniff_format(path: str | Path) -> str:
    """파일 내용(매직바이트·텍스트 패턴)으로 실제 포맷 판별.

    반환: "xls" | "xlsx" | "html" | "csv" | "unknown"
    확장자는 신뢰하지 않는다 — .xls인데 HTML인 파일이 실재한다.
    """
    p = Path(path)
    with open(p, "rb") as f:
        head = f.read(4096)
    if head[:8] == _OLE2_MAGIC:
        return "xls"
    if head[:4] == _ZIP_MAGIC:
        return "xlsx"
    # 텍스트 계열 — BOM 제거 후 판별
    text = head.lstrip(b"\xef\xbb\xbf\xff\xfe\x00 \t\r\n")
    low = text[:512].lower()
    if low.startswith(b"<") and (b"html" in low or b"table" in low or b"?xml" in low):
        return "html"
    # CSV/TSV 추정: 첫 줄에 구분자 존재
    try:
        first_line = text.split(b"\n", 1)[0]
        if (b"," in first_line or b"\t" in first_line) and not low.startswith(b"<"):
            return "csv"
    except Exception:
        pass
    return "unknown"


# ── ③ 금액 파싱 (음수 표기 통일) ──────────────────────────────────────────────

_NEG_PAREN = re.compile(r"^\((.+)\)$")


def parse_amount(val) -> int:
    """회계 표기 금액 문자열 → 정수.

    처리: 쉼표·공백·'원' 제거, float 표기(426584000.0),
          (1,234)→-1234, △1,234/▲1,234→-1234, 끝자리 '-' 부호.
    실패 시 0.
    """
    s = str(val if val is not None else "").strip()
    if not s or s in ("-", "—", "nan", "None"):
        return 0
    # 전각 숫자·기호 → 반각
    s = s.translate(str.maketrans("０１２３４５６７８９－（）", "0123456789-()"))
    s = s.replace(",", "").replace(" ", "").replace("원", "")
    neg = False
    m = _NEG_PAREN.match(s)
    if m:
        neg, s = True, m.group(1)
    if s[:1] in ("△", "▲"):
        neg, s = True, s[1:]
    if s.endswith("-"):          # 일부 ERP의 후행 마이너스 (1234-)
        neg, s = True, s[:-1]
    try:
        n = int(float(s))
    except (ValueError, OverflowError):
        cleaned = re.sub(r"[^\d\-]", "", s)
        if not cleaned or cleaned == "-":
            return 0
        try:
            n = int(cleaned)
        except ValueError:
            return 0
    return -abs(n) if neg else n


# ── ④ 합계행 판별 ─────────────────────────────────────────────────────────────

_TOTAL_EXACT = {"합계", "소계", "총계", "누계", "계", "합  계", "총  계"}


def is_total_row(name: str) -> bool:
    """계정명이 합계·소계행인지 판별 (당기순이익 등 실질 항목은 제외)."""
    n = re.sub(r"[\s\[\]<>()]+", "", str(name))
    if not n:
        return False
    if n in {"합계", "소계", "총계", "누계", "계"}:
        return True
    # 자산총계·부채와자본총계 등 — '...총계'로 끝나면 합계행
    if n.endswith("총계"):
        return True
    # 'Ⅰ.유동자산 계' 류
    if n.endswith("계") and len(n) >= 3 and not n.endswith(("순이익계", "설계", "회계")):
        # '판매비와관리비계' 같은 그룹 소계
        return n[-2] not in ("기", "설", "회")
    return False


# ── 헤더 행 자동 탐지 ─────────────────────────────────────────────────────────

_HEADER_KEYWORDS = {
    "코드", "계정코드", "계정과목", "계정과목명", "계정명", "과목", "과 목",
    "차변", "대변", "금액", "잔액", "기초잔액", "기말잔액",
    "적요", "거래처", "거래처명", "거래처코드",
    "일자", "날짜", "전표일자", "전표번호",
    "당기", "전기", "기초", "기말",
    "자산코드", "자산명", "취득일", "상각률", "상각율",
}


def detect_header_row(df: pd.DataFrame, max_scan: int = 15) -> int | None:
    """제목·빈 행 아래 숨은 실제 헤더 행 인덱스를 찾는다.

    한 행에서 헤더 키워드가 2개 이상 매칭되면 그 행을 헤더로 본다.
    찾지 못하면 None (기존 컬럼 유지).
    """
    for i in range(min(max_scan, len(df))):
        vals = [re.sub(r"\s+", "", str(v)) for v in df.iloc[i].tolist()]
        hits = sum(
            1 for v in vals
            if v and any(kw.replace(" ", "") == v or kw.replace(" ", "") in v
                         for kw in _HEADER_KEYWORDS)
        )
        if hits >= 2:
            return i
    return None


def _apply_header_row(df: pd.DataFrame, row_idx: int) -> pd.DataFrame:
    """row_idx 행을 컬럼명으로 승격하고 그 아래부터 데이터로 사용."""
    new_cols = [str(v).strip() for v in df.iloc[row_idx].tolist()]
    # 빈 컬럼명·중복 컬럼명 처리 (pandas의 .1 규칙과 동일하게)
    seen: dict[str, int] = {}
    final_cols: list[str] = []
    for j, c in enumerate(new_cols):
        name = c if c and c != "nan" else f"Unnamed_{j}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        final_cols.append(name)
    out = df.iloc[row_idx + 1:].copy().reset_index(drop=True)
    out.columns = final_cols
    return out


# ── ⑤ 계정과목명 동의어 (회사·프로그램별 명칭 차이 흡수) ──────────────────────

# 검색 키워드 → 함께 찾아야 할 동의어 (양방향 아님 — 키워드 확장용)
ACCOUNT_SYNONYMS: dict[str, tuple[str, ...]] = {
    "기업업무추진비": ("접대비",),
    "접대비":         ("기업업무추진비",),
    "매출액":         ("영업수익", "매출"),
    "당기순이익":     ("당기순익",),
    "임대보증금":     ("임대보증금예수", "받은보증금", "예수보증금"),
    "단기차입금":     ("당좌차월",),
    "자본총계":       ("자본총액",),
    "퇴직급여충당부채": ("퇴직급여충당금", "퇴직급여충당부채"),
    "복리후생비":     ("복리시설비",),
    "차량유지비":     ("차량비", "차량관리비"),
    "세금과공과":     ("세금과공과금",),
}


def expand_keywords(keywords: tuple[str, ...]) -> tuple[str, ...]:
    """계정명 검색 키워드를 동의어까지 확장한다."""
    out: list[str] = []
    for kw in keywords:
        out.append(kw)
        out.extend(ACCOUNT_SYNONYMS.get(kw, ()))
    # 순서 유지 중복 제거
    return tuple(dict.fromkeys(out))


# ── ② 통합 읽기 (포맷 라우팅 + 정규화 + 진단 메타) ───────────────────────────

_ENCODINGS = ("utf-8", "cp949", "euc-kr", "utf-16")


def _read_html_tables(path: Path) -> tuple[pd.DataFrame, str]:
    """HTML형 파일에서 가장 큰 표를 추출. 반환: (df, 인코딩)."""
    last_err: Exception | None = None
    for enc in _ENCODINGS:
        try:
            with open(path, encoding=enc, errors="strict") as f:
                text = f.read()
            tables = pd.read_html(io.StringIO(text), header=None)
            if not tables:
                continue
            # 셀 수가 가장 많은 표 선택 (머리말·푸터 표 제외 목적)
            df = max(tables, key=lambda t: t.shape[0] * t.shape[1])
            return df.astype(str), enc
        except UnicodeDecodeError as e:
            last_err = e
            continue
        except ValueError as e:   # 표 없음
            last_err = e
            continue
    raise ValueError(f"HTML 표 추출 실패: {last_err}")


def _read_csv_any(path: Path) -> tuple[pd.DataFrame, str]:
    last_err: Exception | None = None
    for enc in _ENCODINGS:
        try:
            df = pd.read_csv(path, dtype=str, encoding=enc,
                             sep=None, engine="python", header=None)
            return df.astype(str), enc
        except (UnicodeDecodeError, pd.errors.ParserError) as e:
            last_err = e
            continue
    raise ValueError(f"CSV 읽기 실패: {last_err}")


def read_any_table(path: str | Path, sheet: str | int = 0) -> tuple[pd.DataFrame, dict]:
    """포맷 자동 판별 → 표 읽기 → 정규화. 반환: (df, 진단 메타).

    정규화 절차:
      1. 포맷 스니핑 (확장자 무시, 내용 기준)
      2. 포맷별 읽기 (HTML·CSV는 인코딩 자동 시도)
      3. 전체 빈 행·빈 열 제거
      4. 헤더 행 자동 탐지 (제목·회사명 행 아래 숨은 실제 헤더)
         — Excel은 1행 헤더가 정상 인식되면 그대로 유지 (기존 동작 보존)
      5. __sheet__/__row__ 감사추적 컬럼 부여

    meta 키: format, encoding, header_row, rows, cols,
             empty_rows_dropped, columns, notes
    """
    p = Path(path)
    fmt = sniff_format(p)
    meta: dict = {
        "format": fmt, "encoding": "", "header_row": 1,
        "rows": 0, "cols": 0, "empty_rows_dropped": 0,
        "columns": [], "notes": [],
    }

    raw_header_offset = 0   # 원본 파일에서 데이터 시작 전 행 수 (감사추적 행번호용)
    if fmt == "html":
        df, enc = _read_html_tables(p)
        meta["encoding"] = enc
        meta["notes"].append("확장자와 달리 실제로는 HTML 저장 파일 — 표 추출로 처리")
        header_idx = detect_header_row(df)
        if header_idx is not None:
            df = _apply_header_row(df, header_idx)
            raw_header_offset = header_idx + 1
            meta["header_row"] = header_idx + 1
        sheet_name = "html"
    elif fmt == "csv":
        df, enc = _read_csv_any(p)
        meta["encoding"] = enc
        header_idx = detect_header_row(df)
        if header_idx is not None:
            df = _apply_header_row(df, header_idx)
            raw_header_offset = header_idx + 1
            meta["header_row"] = header_idx + 1
        sheet_name = "csv"
    else:
        # Excel 계열 — 기존 검증된 동작 유지: header=0으로 읽고,
        # 컬럼 대부분이 Unnamed일 때만 헤더 행 재탐지 (보고서형 보존)
        engine = "xlrd" if fmt == "xls" else "openpyxl"
        fallback = "openpyxl" if engine == "xlrd" else "xlrd"
        df = None
        errs = []
        for eng in (engine, fallback):
            try:
                df = pd.read_excel(p, sheet_name=sheet, header=0,
                                   dtype=str, engine=eng).fillna("")
                meta["encoding"] = eng
                break
            except Exception as e:
                errs.append(f"{eng}: {e}")
        if df is None:
            # 확장자만 xls인 비표준 파일 — HTML 재시도
            try:
                df, enc = _read_html_tables(p)
                meta["format"] = "html"
                meta["encoding"] = enc
                meta["notes"].append("Excel 엔진 실패 → HTML 표 추출로 복구")
                header_idx = detect_header_row(df)
                if header_idx is not None:
                    df = _apply_header_row(df, header_idx)
                    raw_header_offset = header_idx + 1
                    meta["header_row"] = header_idx + 1
            except Exception as e:
                raise ValueError(
                    f"파일 형식을 인식할 수 없습니다: {p.name} ({' | '.join(errs)} | html: {e})"
                )
        else:
            raw_header_offset = 1   # header=0 → 원본 1행이 헤더
            # 헤더 오인식 보정: 컬럼의 60% 이상이 Unnamed면 본문에서 헤더 재탐지
            unnamed = sum(1 for c in df.columns if str(c).startswith("Unnamed"))
            if len(df.columns) > 0 and unnamed / len(df.columns) >= 0.6:
                header_idx = detect_header_row(df)
                if header_idx is not None:
                    df = _apply_header_row(df, header_idx)
                    raw_header_offset += header_idx + 1
                    meta["header_row"] = raw_header_offset
                    meta["notes"].append("제목 행 아래에서 실제 헤더 행을 재탐지")
        if isinstance(sheet, int):
            sheet_name = str(sheet)
            try:
                import openpyxl as _oxl
                wb = _oxl.load_workbook(p, read_only=True, data_only=True)
                if sheet < len(wb.sheetnames):
                    sheet_name = wb.sheetnames[sheet]
                wb.close()
            except Exception:
                pass
        else:
            sheet_name = str(sheet)

    df = df.fillna("")

    # 빈 열 제거 (전체가 빈 문자열인 열) — 단, 이름 있는 컬럼은 유지
    empty_cols = [
        c for c in df.columns
        if str(c).startswith("Unnamed") and not df[c].astype(str).str.strip().any()
    ]
    if empty_cols:
        df = df.drop(columns=empty_cols)

    # 빈 행 제거 + 원본 행번호 보존
    df["__row__"] = range(raw_header_offset + 1, len(df) + raw_header_offset + 1)
    mask_nonempty = df.drop(columns="__row__").astype(str).apply(
        lambda r: any(v.strip() for v in r), axis=1
    )
    meta["empty_rows_dropped"] = int((~mask_nonempty).sum())
    df = df[mask_nonempty].reset_index(drop=True)

    if "__sheet__" not in df.columns:
        df["__sheet__"] = sheet_name

    meta["rows"] = len(df)
    meta["cols"] = len([c for c in df.columns if c not in ("__sheet__", "__row__")])
    meta["columns"] = [str(c) for c in df.columns if c not in ("__sheet__", "__row__")]
    return df, meta
